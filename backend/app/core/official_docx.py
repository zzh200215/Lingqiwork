"""GB/T 9704-2012 党政机关公文版式的 docx 生成——报告模块「交出去」的那一跳。

## 做了什么版式

- 页面 A4，版心按国标：上 37 / 下 35 / 左 28 / 右 26 mm（156×225mm）。
- 标题二号小标宋居中；小节标题三号黑体；正文三号仿宋_GB2312，固定行距 28 磅，
  首行缩进 2 字符（`w:firstLineChars` 字符单位，不是写死的磅值——跟着字号走）。
- `org` 给了才加红头：红色机关标志 + 红色分隔线。**没给就不加**——个人工作台不替
  用户编造机关名。
- 成文日期右空四字（`w:rightChars`）；页码「— N —」宋体四号，**居中**——国标要
  单页码居右、双页码居左，那需要奇偶页分节，第一版先居中，是记录在案的一处偏离。
- 字体名（方正小标宋简体/仿宋_GB2312）照国标写；用户机器上没装时 Word 自行回退。

## Markdown 怎么进来

报告的正文是 Markdown（`## ` 小节、`[n]` 引用、少量加粗与清单）。口径：
- 标题结构由调用方拆好传进来（或走 `parse_markdown`）；
- 行内标记剥掉（加粗/斜体/行内代码/链接取文本）；`[n]` 引用号**照留**——那是证据链；
- 清单：`1.` `1、` 这类数字编号是公文条款的一部分，**原样保留**成独立段；
  `-`/`•` 圆点清单剥掉标记、内容成独立段；
- 空行分段，段内软换行并入同段。
"""
import re
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

# 字号（GB/T 9704-2012：正文三号 = 16pt，标题二号 = 22pt；行距取惯例的固定 28 磅）
BODY_PT, TITLE_PT, LINE_PT = 16, 22, 28
BODY_FONT, HEI_FONT, TITLE_FONT = "仿宋_GB2312", "黑体", "方正小标宋简体"
RED = RGBColor(0xFF, 0x00, 0x00)

# 段内行内 Markdown 的剥离（标题结构不在这里处理，见 parse_markdown）
_INLINE_MD = (
    (re.compile(r"\*\*([^*]+)\*\*"), r"\1"),
    (re.compile(r"(?<!\*)\*([^*\n]+)\*(?!\*)"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
)
# 圆点/短横清单（剥标记）；数字编号不在此列——它是公文条款的一部分
_BULLET_MARKER = re.compile(r"^\s*(?:[-*+•])\s+")
_CODE_FENCE = re.compile(r"^\s*`{3,}")


def parse_markdown(md: str) -> tuple[str, list[dict]]:
    """扁平 Markdown（`# 标题` + `## 小节` + 正文）→ (标题, [{heading, body}])。

    与前端 `stripFrontMatter` 同一条口径：front-matter 是元数据，不进正文。
    没有 `## ` 小节时整篇算一个无标题小节——阅读视图里的散页也要能出。
    """
    text = md
    if text.startswith("---\n"):
        end = text.find("\n---", 3)
        if end >= 0:
            nl = text.find("\n", end + 1)
            text = text[nl + 1 :] if nl >= 0 else ""
    title = ""
    sections: list[dict] = []
    cur: dict | None = None
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("# ") and not s.startswith("##") and not title:
            title = s[2:].strip()
            continue
        if s.startswith("## "):
            cur = {"heading": s[3:].strip(), "body": ""}
            sections.append(cur)
            continue
        if cur is None:
            cur = {"heading": "", "body": ""}
            sections.append(cur)
        cur["body"] += line + "\n"
    # 标题前后的空白会留下「无标题且无正文」的壳——不进小节清单
    sections = [s for s in sections if s["heading"] or s["body"].strip()]
    return title, sections


def _plain(text: str) -> str:
    for pat, rep in _INLINE_MD:
        text = pat.sub(rep, text)
    return text


def _set_font(run, name: str, size: float, *, bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _para(
    doc,
    text: str = "",
    *,
    font: str = BODY_FONT,
    size: float = BODY_PT,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    indent_chars: int = 0,
    right_chars: int = 0,
    line: float = LINE_PT,
    color: RGBColor | None = None,
) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = Pt(line)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    p.alignment = align
    # 缩进用字符单位（firstLineChars/rightChars），Word 按当前字号折算
    if indent_chars:
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), str(indent_chars * 100))
    if right_chars:
        p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:rightChars"), str(right_chars * 100))
    if text:
        _set_font(p.add_run(text), font, size, bold=bold, color=color)


def _red_rule(doc) -> None:
    """红色分隔线：空段落的下边框，粗 3 磅。"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "24")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "FF0000")
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def _page_number_footer(section) -> None:
    """页码「— N —」宋体四号，居中（奇偶分侧是记录在案的偏离，见模块头）。"""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("— "), "宋体", 14)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), "宋体")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "28")  # 半磅单位：28 = 14pt（四号）
    rpr.append(fonts)
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    _set_font(p.add_run(" —"), "宋体", 14)


def _body_paragraphs(doc, body: str) -> None:
    """小节正文 → 公文段落。空行分段；软换行并入同段；数字编号保留为条款。"""
    buf: list[str] = []

    def flush() -> None:
        nonlocal buf
        if buf:
            _para(doc, "".join(buf), indent_chars=2)
            buf = []

    for line in body.splitlines():
        s = _plain(line.strip())
        s = re.sub(r"^&gt;\s?|^>\s?", "", s)  # 引用块标记不进公文正文，内容照排
        if _CODE_FENCE.match(s):
            continue  # 围栏标记本身不进正文；围栏内文字按普通段处理
        if not s:
            flush()
            continue
        if _BULLET_MARKER.match(s):
            flush()
            _para(doc, _BULLET_MARKER.sub("", s), indent_chars=2)
        elif re.match(r"^\s*\d+[.)、]", s):
            flush()
            _para(doc, s, indent_chars=2)  # 数字编号是公文条款，原样保留
        else:
            buf.append(s)
    flush()


def build_docx(*, title: str, sections: list[dict], org: str = "", date_text: str = "") -> bytes:
    """一份报告 → GB/T 9704 版式的 docx 字节。

    `org` 是红头单位名（空 = 不加红头，不编造机关）。`date_text` 空 = 今天。
    """
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(37), Mm(35)
    sec.left_margin, sec.right_margin = Mm(28), Mm(26)

    if org.strip():
        _para(
            doc,
            org.strip(),
            font=TITLE_FONT,
            size=36,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            line=44,
            color=RED,
        )
        _para(doc, "", line=20)
        _red_rule(doc)
        _para(doc, "", line=20)

    # 标题：二号小标宋居中
    _para(doc, title.strip() or "报告", font=TITLE_FONT, size=TITLE_PT, align=WD_ALIGN_PARAGRAPH.CENTER, line=36)

    for s in sections or []:
        heading = (s.get("heading") or "").strip()
        if heading:
            _para(doc, heading, font=HEI_FONT, indent_chars=2)
        _body_paragraphs(doc, s.get("body") or "")

    if not sections:
        _para(doc, "（正文为空）", indent_chars=2)

    _para(doc, "", line=12)
    _para(
        doc,
        date_text.strip() or datetime.now().strftime("%Y年%m月%d日"),
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        right_chars=4,
    )

    _page_number_footer(sec)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
